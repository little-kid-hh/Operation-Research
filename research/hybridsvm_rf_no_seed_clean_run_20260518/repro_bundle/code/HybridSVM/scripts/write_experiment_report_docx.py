# -*- coding: utf-8 -*-
"""
Write a Word (.docx) summary for a HybridSVM experiment folder.

Requires: pip install python-docx

Usage:
  python scripts/write_experiment_report_docx.py \\
    --exp-dir experiments_v2/by_model/qwen3-coder-480b-a35b-instruct/exp_20260413_165559

  python scripts/write_experiment_report_docx.py --exp-dir <path> --output report.docx

  # 旧版 experiments/ 布局：汇总 by_model 下各模型、各次运行的 results.json
  python scripts/write_experiment_report_docx.py --aggregate-v1 [--v1-root experiments] [-o v1_report.docx]
"""

from __future__ import annotations

import argparse
import json
import pathlib
import sys


def _add_complete_pipeline_chapter(doc: object) -> None:
    """
    写入：文档导读 + 四阶段总览 + Stage 1～4 完整中文详述（与 README、evolution.py、evaluate.py 对齐）。
    """
    doc.add_heading("一、文档导读（岛模型说明在哪）", level=1)
    doc.add_paragraph(
        "本 Word 与项目 README 中「Architecture」一节下的「FunSearch 岛模型：进化子流程（详细）」"
        "写的是同一套机制。代码位置：src/evolution.py（Population、Island、EvolutionEngine.step、apply_rules）；"
        "实验默认：run_experiment.py 中的 EVOLUTION_ISLAND_RESET_EVERY（岛重置周期）、HYBRID_RULE_TOP_K（Hybrid 用几条规则）。"
    )

    doc.add_heading("二、四阶段总览", level=1)
    doc.add_paragraph(
        "1. Stage 1 — 线性 SVM：训练或加载模型；对测试集输出 P(y=1) 与硬预测。\n"
        "2. Stage 2 — 难例挖掘：相对阈值 0.5 得 FN/FP/Easy TN/Easy TP；写 hard_case_summary.md。\n"
        "3. Stage 3 — FunSearch 岛模型 + LLM：迭代生成 apply_rule_patch；产出 evolved_rules 与 checkpoint 等。\n"
        "4. Stage 4 — 评估：按全局 top-k 规则链式覆盖得到 Hybrid；写 results.json 与 hard_case_summary_hybrid.md。"
    )

    doc.add_heading("三、Pipeline 分阶段详述（完整）", level=1)

    doc.add_heading("Stage 1：线性 SVM 基线", level=2)
    doc.add_paragraph(
        "在训练集（或配置的 CSV）上拟合线性 SVM（C、核等见 svm_train / run_experiment），"
        "对测试样本输出决策值经 sigmoid 后的 P(y=1)。硬预测一般为 prob≥0.5→1，否则→0。"
        "该基线不参与后续梯度更新；规则只在推理阶段覆盖预测。"
    )

    doc.add_heading("Stage 2：难例挖掘（hard_cases）", level=2)
    doc.add_paragraph(
        "在测试集上对照真标签与 SVM 预测：FN（真可行但 SVM 判否）、FP（真不可行但 SVM 判是）；"
        "在 SVM 判对的样本中划分 Easy TN / Easy TP，用于进化阶段惩罚「乱改判对样本」的行为。"
        "render_hard_case_prompt 等生成 hard_case_summary.md，并可含 easy 典型行与可选对比行（§5），"
        "供 LLM 理解错误模式。实现见 src/hard_cases.py。"
    )

    doc.add_heading("Stage 3：FunSearch 岛模型与 LLM 进化", level=2)
    doc.add_paragraph(
        "【为何多岛】单一列表只保留全局 top 时易早熟收敛；多岛分散存放规则，并周期性对最差岛用强岛好规则克隆重播种，"
        "兼顾探索与利用。日志 diversity_score() 粗测代码多样性。"
    )
    doc.add_paragraph(
        "【结构】Population 含 n_islands 个 Island（默认 4）。每岛多条 Rule（字符串形式的 apply_rule_patch）。"
        "单岛上限 max_island_size（默认 50），超出按 rule_sort_key 截断：score 主序；平局看 FN+FP 总修正；再平局 easy_overrides 越少越好。"
    )
    doc.add_paragraph(
        "【新规则进岛 rule_placement】round_robin（默认）轮询各岛；min_best_score 进当前 best_score 最低的岛；"
        "min_rule_count 进规则条数最少的岛。"
    )
    doc.add_paragraph(
        "【单次 step()】iteration 自增。sample_prompt_rules(k=2)：先随机选一岛，再在该岛内 Boltzmann(temperature=0.5) "
        "采样 k 条作种子。build_evolution_prompt 拼接：特征分位表、SVM 系数摘要、hard_case 全文、k 个种子及统计、RULE_SKELETON；"
        "若 include_prev_iter_in_prompt，附上一轮本批候选最优规则。一次 LLM 调用 → extract_code_from_response 取 apply_rule_patch。"
        "每条候选在 FN/FP/easy TN/easy TP 上 Rule.fit：FN 上应 override→1，FP 上应 override→0，easy 上非 -1 计 easy_overrides；"
        "score 由 FN/FP 两路 F1 风格项平均再乘 easy 惩罚系数。population.add_rule 入岛并可能截断岛大小。"
    )
    doc.add_paragraph(
        "【最差岛重置】当 iteration 能被 island_reset_every 整除时执行 reset_worst_island。"
        "run_experiment 默认 EVOLUTION_ISLAND_RESET_EVERY=5（每 5 步一次；与 EvolutionConfig 类默认 10 不同，以实验入口为准）。"
        "取各岛 best_score 最大者为最优岛、最小者为最差岛；从最优岛取全局最佳 Rule 深拷贝，赋新 id 加入最差岛。"
        "island_reset_mode=trim_top_k：最差岛先只保留分数最高的前 island_reset_keep_top（默认 3）条再追加克隆；"
        "clear_and_clone：先清空最差岛再放入克隆。"
    )
    doc.add_paragraph(
        "【ReEvo】build_reflection_prompt 提供父子规则对比模板，但 EvolutionEngine.step 默认仅一次 LLM，"
        "不另开「反思」第二次调用；跨轮改进主要靠上一轮最优 + 种子规则。"
    )

    doc.add_heading("Stage 4：Hybrid 与评估（apply_rules + evaluate）", level=2)
    doc.add_paragraph(
        "进化结束后汇总所有岛上的 Rule，按 rule_sort_key 取全局前 HYBRID_RULE_TOP_K 条（run_experiment 默认 3），"
        "不是每岛一条。apply_rules：先用 SVM 得初始硬预测；对每个样本按 top-k 列表顺序调用 apply_rule_patch："
        "返回 -1 则试下一条规则；返回 0 或 1 则覆盖并短路停止；全为 -1 则保持 SVM。"
    )
    doc.add_paragraph(
        "evaluate_pipeline 对比 SVM 与 Hybrid 指标；Hybrid 概率在被覆盖处可近似为 0.95/0.05（见 evaluate.py），"
        "故 Hybrid 的 AUC 等需结合场景解读。results.json 记录总体指标；hard_case_summary_hybrid 在 Hybrid 预测上再挖 FN/FP。"
    )


def _load_json(path: pathlib.Path) -> dict | None:
    if not path.is_file():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def _count_rules(evolved_path: pathlib.Path) -> int | None:
    if not evolved_path.is_file():
        return None
    data = json.loads(evolved_path.read_text(encoding="utf-8"))
    return len(data) if isinstance(data, list) else None


def _model_display_name(exp_dir: pathlib.Path, results: dict) -> str:
    lm = results.get("llm_model")
    if lm:
        return str(lm)
    parts = exp_dir.parts
    if "by_model" in parts:
        i = parts.index("by_model")
        if i + 1 < len(parts):
            slug = parts[i + 1]
            if slug == "_legacy_unknown":
                return "_legacy_unknown（results 中无 llm_model，以目录为准）"
            return slug
    return "—"


def _discover_v1_runs(legacy_root: pathlib.Path) -> list[pathlib.Path]:
    """返回含 results.json 的实验目录列表（experiments/by_model/<slug>/exp_*/）。"""
    by_model = legacy_root / "by_model"
    if not by_model.is_dir():
        return []
    dirs: list[pathlib.Path] = []
    for p in sorted(by_model.rglob("results.json")):
        parent = p.parent
        if parent.name.startswith("exp_") and (parent / "results.json").is_file():
            dirs.append(parent)
    # 去重并保持顺序（rglob 已排序）
    seen: set[pathlib.Path] = set()
    unique: list[pathlib.Path] = []
    for d in dirs:
        r = d.resolve()
        if r not in seen:
            seen.add(r)
            unique.append(r)
    return unique


def _add_run_metrics_and_notes(
    doc: object,
    results: dict | None,
    *,
    include_interpretation: bool = True,
) -> None:
    """在文档中追加单次运行的指标表、混淆矩阵、难例规模（与 build_document 一致）。"""
    if not results:
        doc.add_paragraph("未找到或无法解析 results.json。")
        return

    svm = results.get("svm", {})
    hyb = results.get("hybrid", {})
    skip_evolution = bool(results.get("skip_evolution", False))

    if skip_evolution:
        doc.add_paragraph(
            "说明：本次为 baseline-only 运行，未执行 evolution。为统一结果结构，`hybrid` 字段按 baseline 原样回填。"
        )

    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "指标"
    hdr[1].text = "SVM"
    hdr[2].text = "Hybrid"
    hdr[3].text = "差值 (Hybrid − SVM)"

    metrics = [
        ("Accuracy", "accuracy"),
        ("Precision", "precision"),
        ("Recall", "recall"),
        ("ROC AUC", "auc"),
        ("FPR", "fpr"),
        ("FNR", "fnr"),
        ("TPR@FPR=1%", "tpr_at_fpr1pct"),
    ]
    for label, key in metrics:
        row = table.add_row().cells
        sv = svm.get(key)
        hv = hyb.get(key)
        row[0].text = label
        row[1].text = f"{float(sv):.4f}" if sv is not None else "—"
        row[2].text = f"{float(hv):.4f}" if hv is not None else "—"
        if sv is not None and hv is not None:
            row[3].text = f"{float(hv) - float(sv):+.4f}"
        else:
            row[3].text = "—"

    doc.add_paragraph("")
    doc.add_paragraph("混淆矩阵（TN / FP / FN / TP）")

    t2 = doc.add_table(rows=3, cols=3)
    t2.rows[0].cells[0].text = ""
    t2.rows[0].cells[1].text = "SVM"
    t2.rows[0].cells[2].text = "Hybrid"
    t2.rows[1].cells[0].text = "TN / FP / FN / TP"
    t2.rows[1].cells[1].text = (
        f"{svm.get('tn')}/{svm.get('fp')}/{svm.get('fn')}/{svm.get('tp')}"
    )
    t2.rows[1].cells[2].text = (
        f"{hyb.get('tn')}/{hyb.get('fp')}/{hyb.get('fn')}/{hyb.get('tp')}"
    )

    doc.add_heading("难例规模", level=3)
    nh = results.get("n_hard_cases", {})
    doc.add_paragraph(
        f"SVM 视角：FN={nh.get('fn')}, FP={nh.get('fp')}, "
        f"Easy TN={nh.get('easy_tn')}, Easy TP={nh.get('easy_tp')}"
    )
    nhh = results.get("n_hard_cases_hybrid", {})
    if nhh:
        doc.add_paragraph(
            f"最终 Hybrid 再挖掘（相对 Hybrid 预测）：FN={nhh.get('fn')}, FP={nhh.get('fp')}"
        )

    if include_interpretation:
        doc.add_heading("解读说明", level=3)
        doc.add_paragraph(
            "Hybrid 在全体准确率上常低于纯 SVM，因规则会在大量「原判对」样本上触发覆盖；"
            "难例子集（原 SVM 的 FN+FP）上的修正率更反映规则价值。"
            "详见项目 README 与设计原则。"
        )


def build_aggregate_v1_document(legacy_root: pathlib.Path) -> "Document":
    """旧版 experiments/by_model/ 下所有 exp_* 运行的汇总报告（中文）。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    legacy_root = legacy_root.resolve()
    runs = _discover_v1_runs(legacy_root)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("HybridSVM 旧版布局实验汇总（v1）", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        f"根目录：{legacy_root}。本文档由 scripts/write_experiment_report_docx.py "
        f"--aggregate-v1 扫描 by_model 下各 exp_* 目录的 results.json 自动生成。"
    )

    _add_complete_pipeline_chapter(doc)

    doc.add_heading("四、各模型 / 各次运行总览", level=1)
    if not runs:
        doc.add_paragraph("未在 by_model/*/exp_*/ 下发现 results.json。")
        return doc

    loaded: list[tuple[pathlib.Path, dict | None]] = []
    for exp_dir in runs:
        loaded.append((exp_dir, _load_json(exp_dir / "results.json")))

    sum_table = doc.add_table(rows=1, cols=8)
    sh = sum_table.rows[0].cells
    sh[0].text = "模型标识"
    sh[1].text = "实验目录名"
    sh[2].text = "时间戳"
    sh[3].text = "续跑"
    sh[4].text = "SVM Acc"
    sh[5].text = "Hybrid Acc"
    sh[6].text = "Δ Acc"
    sh[7].text = "best_rule_score"

    for exp_dir, results in loaded:
        row = sum_table.add_row().cells
        r = results or {}
        svm = r.get("svm", {})
        hyb = r.get("hybrid", {})
        sa = svm.get("accuracy")
        ha = hyb.get("accuracy")
        br = hyb.get("best_rule_score")
        row[0].text = _model_display_name(exp_dir, r)
        row[1].text = exp_dir.name
        row[2].text = str(r.get("timestamp", "—"))
        if r.get("resumed"):
            row[3].text = f"是 (from {r.get('resume_start_iteration', '?')})"
        else:
            row[3].text = "否"
        row[4].text = f"{float(sa):.4f}" if sa is not None else "—"
        row[5].text = f"{float(ha):.4f}" if ha is not None else "—"
        if sa is not None and ha is not None:
            row[6].text = f"{float(ha) - float(sa):+.4f}"
        else:
            row[6].text = "—"
        row[7].text = f"{float(br):.6f}" if br is not None else "—"

    doc.add_paragraph("")
    doc.add_paragraph(
        "说明：同一模型目录下可能有多次运行（不同 exp_* 文件夹）；"
        "_legacy_unknown 表示早期运行未在 results.json 中记录 llm_model。"
    )

    doc.add_heading("五、分运行明细（按目录排序）", level=1)
    for idx, (exp_dir, results) in enumerate(loaded, start=1):
        if not results:
            doc.add_heading(f"{idx}. {exp_dir.name}（无 results）", level=2)
            doc.add_paragraph(str(exp_dir))
            continue

        name = _model_display_name(exp_dir, results)
        doc.add_heading(f"{idx}. {name} — {exp_dir.name}", level=2)

        p = doc.add_paragraph()
        p.add_run("完整路径：").bold = True
        p.add_run(str(exp_dir))

        doc.add_heading("运行信息", level=3)
        doc.add_paragraph(f"时间戳：{results.get('timestamp', '—')}")
        doc.add_paragraph(f"LLM 模型（字段）：{results.get('llm_model', '（未记录）')}")
        doc.add_paragraph(f"进化迭代步数：{results.get('n_evolution_iters', '—')}")
        if results.get("resumed"):
            doc.add_paragraph(
                f"续跑：是（自 iteration {results.get('resume_start_iteration', '?')} 起）"
            )
        else:
            doc.add_paragraph("续跑：否")

        evolved_path = exp_dir / "evolved_rules.json"
        n_rules = _count_rules(evolved_path)
        if n_rules is not None:
            doc.add_paragraph(f"最终 evolved 规则条数：{n_rules}")
        br = results.get("hybrid", {}).get("best_rule_score")
        if br is not None:
            doc.add_paragraph(f"最优规则分数（checkpoint 记录）：{float(br):.6f}")

        doc.add_heading("测试集指标：SVM vs Hybrid", level=3)
        _add_run_metrics_and_notes(doc, results, include_interpretation=False)

    doc.add_heading("六、整体解读说明", level=1)
    doc.add_paragraph(
        "Hybrid 在全体准确率上常低于纯 SVM，因规则会在大量「原判对」样本上触发覆盖；"
        "难例子集（原 SVM 的 FN+FP）上的修正率更反映规则价值。"
        "详见项目 README 与设计原则。"
    )

    return doc


def build_document(exp_dir: pathlib.Path) -> "Document":
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    exp_dir = exp_dir.resolve()
    results_path = exp_dir / "results.json"
    results = _load_json(results_path)
    evolved_path = exp_dir / "evolved_rules.json"
    n_rules = _count_rules(evolved_path)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("HybridSVM 实验报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "本文档由 scripts/write_experiment_report_docx.py 从 results.json 等自动生成。"
    )

    _add_complete_pipeline_chapter(doc)

    doc.add_heading("四、本次运行信息", level=1)
    p = doc.add_paragraph()
    p.add_run("实验目录：").bold = True
    p.add_run(str(exp_dir))
    if results:
        doc.add_paragraph(f"时间戳：{results.get('timestamp', '—')}")
        doc.add_paragraph(f"LLM 模型：{results.get('llm_model', '—')}")
        doc.add_paragraph(f"进化迭代步数：{results.get('n_evolution_iters', '—')}")
        if results.get("resumed"):
            doc.add_paragraph(
                f"续跑：是（自 iteration {results.get('resume_start_iteration', '?')} 起）"
            )
        else:
            doc.add_paragraph("续跑：否")
        if results.get("skip_evolution"):
            doc.add_paragraph("是否跳过进化：是（baseline-only）")
        if n_rules is not None:
            doc.add_paragraph(f"最终 evolved 规则条数：{n_rules}")
        br = results.get("hybrid", {}).get("best_rule_score")
        if br is not None:
            doc.add_paragraph(f"最优规则分数（checkpoint 记录）：{float(br):.6f}")

    doc.add_heading("五、测试集指标：SVM vs Hybrid", level=1)
    if not results:
        doc.add_paragraph("未找到 results.json，无法填写指标表。")
        return doc

    _add_run_metrics_and_notes(doc, results, include_interpretation=False)

    doc.add_heading("六、解读说明", level=1)
    doc.add_paragraph(
        "Hybrid 在全体准确率上常低于纯 SVM，因规则会在大量「原判对」样本上触发覆盖；"
        "难例子集（原 SVM 的 FN+FP）上的修正率更反映规则价值。"
        "详见项目 README 与设计原则。"
    )

    return doc


def main() -> None:
    parser = argparse.ArgumentParser(description="Write HybridSVM experiment summary to .docx")
    grp = parser.add_mutually_exclusive_group(required=True)
    grp.add_argument(
        "--exp-dir",
        type=pathlib.Path,
        help="Experiment folder containing results.json (e.g. experiments_v2/.../exp_TIMESTAMP)",
    )
    grp.add_argument(
        "--aggregate-v1",
        action="store_true",
        help="Scan legacy experiments/by_model/*/exp_*/results.json and write one combined report",
    )
    parser.add_argument(
        "--v1-root",
        type=pathlib.Path,
        default=pathlib.Path("experiments"),
        help="Root for --aggregate-v1 (default: experiments)",
    )
    parser.add_argument(
        "--output",
        "-o",
        type=pathlib.Path,
        default=None,
        help="Output .docx path (default: <exp-dir>/experiment_report.docx or <v1-root>/v1_all_models_experiment_report.docx)",
    )
    args = parser.parse_args()

    try:
        from docx import Document  # noqa: F401
    except ImportError:
        print(
            "ERROR: python-docx is not installed. Run: pip install python-docx",
            file=sys.stderr,
        )
        sys.exit(1)

    if args.aggregate_v1:
        legacy_root = args.v1_root.resolve()
        if not legacy_root.is_dir():
            print(f"ERROR: not a directory: {legacy_root}", file=sys.stderr)
            sys.exit(1)
        out = args.output or (legacy_root / "v1_all_models_experiment_report.docx")
        out = out.resolve()
        doc = build_aggregate_v1_document(legacy_root)
        doc.save(str(out))
        print(f"Wrote: {out}")
        return

    exp_dir = args.exp_dir.resolve()
    if not exp_dir.is_dir():
        print(f"ERROR: not a directory: {exp_dir}", file=sys.stderr)
        sys.exit(1)

    out = args.output or (exp_dir / "experiment_report.docx")
    out = out.resolve()
    doc = build_document(exp_dir)
    doc.save(str(out))
    print(f"Wrote: {out}")


if __name__ == "__main__":
    main()
